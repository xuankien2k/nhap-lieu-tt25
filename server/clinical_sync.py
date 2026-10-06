"""Đồng bộ khám lâm sàng từ file có mã KCB, khớp người trong đoàn."""

from datetime import date, datetime

from openpyxl.utils.datetime import from_excel

from server.excel_io import load_workbook_bytes, parse_date
from server.kcb_sync import code_text, fold_name

NOI_PAIRS = [
    ("X", "Y"),
    ("Z", "AA"),
    ("AB", "AC"),
    ("AD", "AE"),
    ("AF", "AG"),
    ("AH", "AI"),
    ("AJ", "AK"),
    ("AL", "AM"),
]
SPECIALTIES = [
    ("noi khoa", None, None),
    ("mat", "AV", "AW"),
    ("tai mui hong", "BC", "BD"),
    ("rang ham mat", "BH", "BI"),
    ("da lieu", "BK", "BL"),
    ("phu san", "BN", "BO"),
    ("ngoai khoa", "AO", "AP"),
]
CLINICAL_COLUMNS = [
    "O", "P", "Q", "R", "S", "T", "W",
    "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM",
    "AO", "AP", "AV", "AW", "BC", "BD", "BH", "BI", "BK", "BL", "BN", "BO",
    "CW", "CX",
    "REVIEW_COLS",
]

ROMAN_GRADE = {"I": "1", "II": "2", "III": "3", "IV": "4", "V": "5", "1": "1", "2": "2", "3": "3", "4": "4", "5": "5"}


def header_key(value):
    text = str(value or "").replace("-", " ").replace("/", " ")
    return fold_name(text)


def number_text(value):
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip().replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return text
    return str(int(number)) if number.is_integer() else f"{number:g}"


def date_text(value):
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).strftime("%d/%m/%Y")
        except Exception:
            return ""
    text = str(value or "").strip()
    if not text:
        return ""
    parsed = parse_date(text[:10] if " " in text else text)
    return parsed.strftime("%d/%m/%Y") if parsed else text


def split_bp(value):
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return "", ""
    for sep in ("/", "-", "\\"):
        if sep in text:
            systolic, diastolic = text.split(sep, 1)
            return number_text(systolic), number_text(diastolic)
    return number_text(text), ""


def classify(value):
    text = " ".join(str(value or "").replace("\n", " ").split())
    key = fold_name(text)
    if not key:
        return "empty", ""
    if key in {"-", "—", "khong", "khong co"}:
        return "na", text
    if key in {"binh thuong", "bt", "binh thuong."}:
        return "normal", "Bình thường"
    return "abnormal", text


def _pick(labels, *names):
    for name in names:
        if name in labels:
            return labels[name]
    return None


def _sheet_columns(ws):
    for row in range(1, 8):
        labels = {}
        for scan in range(row, min(row + 3, 10)):
            for col in range(1, min(ws.max_column, 80) + 1):
                raw = ws.cell(scan, col).value
                if raw is None or len(str(raw)) > 40:
                    continue
                key = header_key(raw)
                if key:
                    labels.setdefault(key, col)
        code_col = _pick(labels, "ma kcb")
        bp_col = _pick(labels, "huyet ap")
        if code_col and bp_col:
            start = row + 1
            for data_row in range(row + 1, row + 6):
                if code_text(ws.cell(data_row, code_col).value).isdigit():
                    start = data_row
                    break
            return labels, start
    return None, None


def fields_from_row(ws, row, columns):
    fields = {col: "" for col in CLINICAL_COLUMNS if col != "REVIEW_COLS"}
    review = []
    exam_date = date_text(ws.cell(row, columns["ngay kham"]).value) if columns.get("ngay kham") else ""
    if exam_date:
        fields["O"] = exam_date
        fields["W"] = exam_date
    if columns.get("can nang"):
        fields["P"] = number_text(ws.cell(row, columns["can nang"]).value)
    if columns.get("chieu cao"):
        fields["Q"] = number_text(ws.cell(row, columns["chieu cao"]).value)
    if columns.get("mach"):
        fields["R"] = number_text(ws.cell(row, columns["mach"]).value)
    if columns.get("huyet ap"):
        systolic, diastolic = split_bp(ws.cell(row, columns["huyet ap"]).value)
        fields["S"] = systolic
        fields["T"] = diastolic

    noi_col = columns.get("noi khoa")
    if noi_col:
        kind, text = classify(ws.cell(row, noi_col).value)
        if kind == "normal":
            for text_col, grade_col in NOI_PAIRS:
                fields[text_col] = text
                fields[grade_col] = "1"
        elif kind == "abnormal":
            fields["X"] = text
            review.append("X")
            for _text_col, grade_col in NOI_PAIRS:
                review.append(grade_col)
        elif kind == "na":
            fields["X"] = text

    for key, text_col, grade_col in SPECIALTIES:
        if key == "noi khoa" or not columns.get(key):
            continue
        kind, text = classify(ws.cell(row, columns[key]).value)
        fields[text_col] = text
        if kind == "normal":
            fields[grade_col] = "1"
        elif kind == "abnormal":
            review.extend([text_col, grade_col])
    if columns.get("phan loai suc khoe"):
        grade = str(ws.cell(row, columns["phan loai suc khoe"]).value or "").strip().upper()
        fields["CW"] = ROMAN_GRADE.get(grade, "")
    if columns.get("ket luan"):
        fields["CX"] = str(ws.cell(row, columns["ket luan"]).value or "").strip()
    fields["REVIEW_COLS"] = ",".join(review)
    return fields


def parse_clinical_workbook(raw):
    wb = load_workbook_bytes(raw, data_only=True)
    ws = None
    columns = None
    start = None
    best_count = -1
    for name in wb.sheetnames:
        found, data_start = _sheet_columns(wb[name])
        if not found:
            continue
        code_col = _pick(found, "ma kcb")
        count = 0
        for row in range(data_start, wb[name].max_row + 1):
            if code_text(wb[name].cell(row, code_col).value).isdigit():
                count += 1
        if count > best_count:
            best_count = count
            ws = wb[name]
            start = data_start
            columns = {
                "ma kcb": _pick(found, "ma kcb"),
                "ngay kham": _pick(found, "ngay kham"),
                "chieu cao": _pick(found, "chieu cao"),
                "can nang": _pick(found, "can nang"),
                "mach": _pick(found, "mach"),
                "huyet ap": _pick(found, "huyet ap"),
                "noi khoa": _pick(found, "noi khoa"),
                "mat": _pick(found, "mat"),
                "tai mui hong": _pick(found, "tai mui hong"),
                "rang ham mat": _pick(found, "rang ham mat"),
                "da lieu": _pick(found, "da lieu"),
                "phu san": _pick(found, "phu san"),
                "ngoai khoa": _pick(found, "ngoai khoa"),
                "phan loai suc khoe": _pick(found, "phan loai suc khoe"),
                "ket luan": _pick(found, "ket luan"),
            }
    if ws is None:
        raise ValueError("File cần có cột Mã KCB và Huyết áp.")
    results = {}
    for row in range(start, ws.max_row + 1):
        code = code_text(ws.cell(row, columns["ma kcb"]).value)
        if not code or not code.isdigit():
            continue
        results[code] = fields_from_row(ws, row, columns)
    if not results:
        raise ValueError("Không thấy dòng khám lâm sàng nào có mã KCB.")
    return results


def plan_clinical_sync(patients, results):
    with_code = [patient for patient in patients if (patient["data"].get("MA_KCB") or "").strip()]
    if not with_code:
        raise ValueError("Chưa đồng bộ mã KCB. Hãy đồng bộ mã KCB trước khi đồng bộ khám lâm sàng.")
    by_code = {patient["data"]["MA_KCB"].strip() for patient in with_code}
    updates = {}
    missing = []
    for patient in with_code:
        code = patient["data"]["MA_KCB"].strip()
        fields = results.get(code)
        if not fields:
            missing.append({"name": patient["data"].get("B") or "Chưa có tên", "ma_kcb": code})
            continue
        fields = dict(fields)
        if str(patient["data"].get("D") or "") == "1":
            fields["BN"] = ""
            fields["BO"] = ""
            fields["BP"] = ""
            fields["REVIEW_COLS"] = ",".join(
                col for col in fields.get("REVIEW_COLS", "").split(",") if col and col not in {"BN", "BO", "BP"}
            )
        updates[patient["id"]] = fields
    return {
        "updates": updates,
        "missing_in_file": missing,
        "skipped_not_in_batch": len(set(results) - by_code),
        "without_code": [
            {"name": patient["data"].get("B") or "Chưa có tên"}
            for patient in patients
            if not (patient["data"].get("MA_KCB") or "").strip()
        ],
    }
