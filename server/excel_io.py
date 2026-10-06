import io
import re
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import from_excel

from server.schema import COLUMNS, FIELD_MAP

warnings.filterwarnings("ignore", message="Data Validation extension is not supported")

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "template" / "mau-ksk.xlsx"
SHEET_NAME = "Ket qua KSKDK"


def _sheet(wb):
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    return wb.active


def _format_float(value):
    text = f"{value:.10f}".rstrip("0").rstrip(".")
    return text


def cell_to_text(cell, field):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if field["type"] == "date" and isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).strftime("%d/%m/%Y")
        except Exception:
            return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return _format_float(value)
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def load_workbook_bytes(raw, data_only=True):
    try:
        return openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only)
    except Exception:
        fixed = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(fixed, "w") as zout:
            for item in zin.infolist():
                zout.writestr(item.filename.replace("\\", "/"), zin.read(item.filename))
        fixed.seek(0)
        return openpyxl.load_workbook(fixed, data_only=data_only)


def parse_workbook(raw):
    wb = load_workbook_bytes(raw, data_only=False)
    ws = _sheet(wb)
    header = ws["B4"].value
    if not header or "Họ" not in str(header):
        raise ValueError("File không đúng mẫu sổ khám sức khỏe. Cần cột Họ và tên ở dòng tiêu đề.")
    records = []
    for row in range(6, ws.max_row + 1):
        data = {}
        filled = False
        for field in FIELD_MAP.values():
            text = cell_to_text(ws[f"{field['col']}{row}"], field)
            data[field["col"]] = text
            if text:
                filled = True
        if filled:
            records.append(data)
    return records


def parse_date(value):
    text = value.strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def parse_number(value):
    text = value.strip().replace(" ", "").replace(",", ".")
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        number = float(text)
        return int(number) if number.is_integer() else number
    return None


def _write_cell(cell, field, raw):
    text = "" if raw is None else str(raw).strip()
    if text == "":
        cell.value = None
        return
    if field["type"] == "date":
        cell.value = parse_date(text) or text
        return
    if field["type"] in ("number", "grade"):
        number = parse_number(text)
        cell.value = number if number is not None else text
        return
    cell.value = text


def ensure_template():
    if TEMPLATE.is_file():
        return
    source = Path("/Users/xuankien/Downloads/HAUGIANG.xlsx")
    if not source.is_file():
        raise FileNotFoundError("Thiếu file mẫu template/mau-ksk.xlsx")
    wb = openpyxl.load_workbook(source)
    ws = _sheet(wb)
    for row in ws.iter_rows(min_row=6, max_row=max(ws.max_row, 6), max_col=107):
        for cell in row:
            if cell.value is not None:
                cell.value = None
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TEMPLATE)


def export_workbook(source_path, patients):
    path = Path(source_path) if source_path else TEMPLATE
    if not path.is_file():
        path = TEMPLATE
    wb = openpyxl.load_workbook(path)
    ws = _sheet(wb)
    last = max(ws.max_row, 5 + len(patients))
    for row in ws.iter_rows(min_row=6, max_row=last, max_col=107):
        for cell in row:
            if cell.value is not None:
                cell.value = None
    for index, patient in enumerate(patients):
        row = 6 + index
        data = patient["data"] if isinstance(patient, dict) and "data" in patient else patient
        for col in COLUMNS:
            _write_cell(ws[f"{col}{row}"], FIELD_MAP[col], data.get(col, ""))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def blank_record():
    return {col: "" for col in COLUMNS}
