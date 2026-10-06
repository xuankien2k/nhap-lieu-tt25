"""Gắn mã KCB từ danh sách bệnh nhân (cột Tên BN, Ngày sinh, Mã KCB)."""

import unicodedata
from collections import defaultdict
from datetime import date, datetime

from openpyxl.utils.datetime import from_excel

from server.excel_io import load_workbook_bytes, parse_date


def fold_name(value):
    text = " ".join(str(value or "").split())
    decomposed = unicodedata.normalize("NFD", text)
    stripped = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    return stripped.replace("đ", "d").replace("Đ", "D").casefold()


def dob_text(value):
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).strftime("%d/%m/%Y")
        except Exception:
            return ""
    text = str(value or "").strip()
    if not text:
        return ""
    head = text[:10] if " " in text else text
    parsed = parse_date(head) or parse_date(text)
    return parsed.strftime("%d/%m/%Y") if parsed else text


def code_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _header_map(ws):
    for row in range(1, min(ws.max_row, 15) + 1):
        found = {}
        for col in range(1, min(ws.max_column, 40) + 1):
            label = fold_name(ws.cell(row, col).value)
            if label == "ma kcb":
                found["ma_kcb"] = col
            elif label in ("ten bn", "ho va ten", "ho ten"):
                found["name"] = col
            elif label in ("ngay sinh", "ngay thang nam sinh"):
                found["dob"] = col
        if {"ma_kcb", "name", "dob"} <= found.keys():
            return row, found
    return None, None


def parse_kcb_list(raw):
    wb = load_workbook_bytes(raw, data_only=True)
    ws = wb.active
    header_row, columns = _header_map(ws)
    if not columns:
        raise ValueError("File cần có cột Tên BN, Ngày sinh và Mã KCB.")
    rows = []
    for index in range(header_row + 1, ws.max_row + 1):
        name = " ".join(str(ws.cell(index, columns["name"]).value or "").split())
        code = code_text(ws.cell(index, columns["ma_kcb"]).value)
        dob = dob_text(ws.cell(index, columns["dob"]).value)
        if not name and not code:
            continue
        if not name or not code:
            continue
        rows.append({"name": name, "dob": dob, "ma_kcb": code})
    if not rows:
        raise ValueError("Không thấy dòng bệnh nhân nào có mã KCB.")
    return rows


def plan_sync(patients, rows):
    exact = defaultdict(list)
    by_name = defaultdict(list)
    for index, row in enumerate(rows):
        exact[(fold_name(row["name"]), row["dob"])].append(index)
        by_name[fold_name(row["name"])].append(index)

    patients_by_name = defaultdict(list)
    for patient in patients:
        patients_by_name[fold_name(patient["data"].get("B"))].append(patient)

    used = set()
    assignments = {}
    name_only = []
    unmatched_patients = []
    ambiguous = []

    for patient in patients:
        display = patient["data"].get("B") or "Chưa có tên"
        dob = (patient["data"].get("C") or "").strip()
        name = fold_name(display)
        hits = exact.get((name, dob), [])
        kind = "exact"
        if len(hits) > 1:
            ambiguous.append({"name": display, "dob": dob, "reason": "Nhiều dòng trong file trùng họ tên và ngày sinh"})
            continue
        if len(hits) == 1:
            index = hits[0]
        else:
            name_hits = by_name.get(name, [])
            if len(name_hits) == 1 and len(patients_by_name.get(name, [])) == 1:
                index = name_hits[0]
                kind = "name"
            elif len(name_hits) > 1:
                ambiguous.append({"name": display, "dob": dob, "reason": "Trùng tên, ngày sinh không khớp"})
                continue
            else:
                unmatched_patients.append({"name": display, "dob": dob})
                continue
        if index in used:
            ambiguous.append({"name": display, "dob": dob, "reason": "Dòng trong file đã gắn cho người khác"})
            continue
        used.add(index)
        code = rows[index]["ma_kcb"]
        assignments[patient["id"]] = code
        if kind == "name":
            name_only.append(
                {
                    "name": display,
                    "patient_dob": dob,
                    "file_dob": rows[index]["dob"],
                    "ma_kcb": code,
                }
            )

    unmatched_file = [
        {"name": rows[index]["name"], "dob": rows[index]["dob"], "ma_kcb": rows[index]["ma_kcb"]}
        for index in range(len(rows))
        if index not in used
    ]
    return {
        "assignments": assignments,
        "name_only": name_only,
        "unmatched_patients": unmatched_patients,
        "unmatched_file": unmatched_file,
        "ambiguous": ambiguous,
    }
